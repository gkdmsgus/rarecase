# Clinical Checkpoint 논문을 학회 양식(2쪽, 2단)으로 docx 생성
# 실행: python tools/build-paper-docx.py
# 양식 근거: 논문_CMM_검토_20251016.pdf (한국실천공학교육학회 학술발표대회 논문집)
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "docs", "figures")
OUT = os.path.join(ROOT, "docs", "Clinical_Checkpoint_paper.docx")

SANS = "맑은 고딕"
SERIF = "바탕"


def set_font(run, name=SERIF, size=9.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(doc, text="", *, style_size=9.0, bold=False, align=None, font=SERIF,
         space_after=2, space_before=0, indent_first=True, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.06
    if align is not None:
        p.alignment = align
    if indent_first and align is None:
        pf.first_line_indent = Cm(0.4)
    if text:
        set_font(p.add_run(text), font, style_size, bold, color)
    return p


def columns(section, num):
    """구역의 단 수를 지정한다 (python-docx가 직접 지원하지 않아 XML로 설정)."""
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), "425")  # 0.75cm


def figure(doc, filename, caption, width_cm=7.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(os.path.join(FIG, filename), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    set_font(cap.add_run(caption), SANS, 8.5)


def heading(doc, text):
    para(doc, text, style_size=11, bold=True, font=SANS,
         space_before=8, space_after=3, indent_first=False)


doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.0)
s.top_margin = s.bottom_margin = Cm(2.0)

# ── 제목 블록 (1단) ───────────────────────────────────────────
para(doc, "생성형 AI 기반 가상환자 문진과 근거보고서형 디브리핑을 결합한",
     style_size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=SERIF, space_after=0)
para(doc, "임상추론 교육 시뮬레이터 설계", style_size=15, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, font=SERIF, space_after=4)
para(doc, "[저자명 입력]", style_size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para(doc, "[소속 입력]", style_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para(doc, "[저자별 이메일 입력]", style_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=7)

para(doc, "Design of a Clinical Reasoning Education Simulator Combining",
     style_size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para(doc, "Generative AI-Based Virtual Patient Interviews and Evidence-Report Debriefing",
     style_size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para(doc, "[English author names]", style_size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para(doc, "[English affiliation]", style_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=7)

para(doc, "요    약", style_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
     font=SANS, space_after=4)
para(doc,
     "본 연구는 의학 및 보건계열 학습자가 고위험 임상증례를 반복 학습할 수 있도록, 생성형 AI 기반 가상환자 문진, "
     "신체 진찰, 교육용 영상 관찰 및 근거보고서형 디브리핑을 결합한 웹 기반 임상추론 교육 시뮬레이터를 제안한다. "
     "학습자는 자유 문장으로 질문하고, 증례의 전문가 검수 정보 범위 안에서 가상환자의 응답을 받는다. 이후 관찰 내용과 "
     "검사 선택 이유를 보고서로 제출하며, 전문가 검토 기준을 통해 누락 단서와 불확실성을 확인한다. AI는 진단 정답을 "
     "판정하거나 의료 결정을 대신하지 않고, 자연어 상호작용과 학습 근거의 구조화에 한정하여 활용된다.",
     style_size=9.0, space_after=6)
para(doc, "Key Words : 생성형 AI, 가상환자, 임상추론, 교육용 의료영상, 디브리핑",
     style_size=9.0, font=SANS, indent_first=False, space_after=6)

# ── 본문 (2단) ────────────────────────────────────────────────
body = doc.add_section(WD_SECTION.CONTINUOUS)
body.left_margin = body.right_margin = Cm(2.0)
body.top_margin = body.bottom_margin = Cm(2.0)
columns(body, 2)

heading(doc, "Ⅰ. 서 론")
para(doc,
     "임상교육에서 문진과 임상추론은 정답을 고르는 일이 아니라, 환자의 발화에서 위험 단서를 수집하고 관찰 사실과 "
     "판단 근거를 연결하는 과정이다. 그러나 희귀하거나 시간 의존성이 높은 응급 증례는 실제 현장에서 반복 경험하기 "
     "어렵고, 기존의 선택지 중심 디지털 콘텐츠는 학습자가 질문을 구성하고 근거를 설명하는 과정을 충분히 드러내기 "
     "어렵다. 이에 본 연구는 자유 문진과 근거 제출을 중심으로 한 가상환자 기반 학습 환경을 설계하였다.")

heading(doc, "Ⅱ. 본 론")
para(doc, "1. 시스템 설계", style_size=10, bold=True, font=SANS,
     space_before=5, space_after=2, indent_first=False)
para(doc,
     "Clinical Checkpoint는 문진, 신체 진찰, 검사 선택, 영상 관찰, 학습 보고서 제출의 순서로 구성된다. 초기 "
     "시나리오는 급성 비외상성 두통 상황을 대상으로 한다. 학습자는 가상환자에게 필요한 질문을 직접 입력하고, AI는 "
     "전문가가 검수한 증례 사실을 벗어나지 않는 범위에서 환자 발화를 생성한다. 즉 AI는 진단이나 통과 여부를 "
     "판정하지 않고, 자연스러운 상호작용을 구현한다.")
figure(doc, "architecture.png", "그림 1. 시스템 구성과 생성형 AI의 역할 범위", width_cm=7.2)

para(doc, "2. 학습 흐름", style_size=10, bold=True, font=SANS,
     space_before=5, space_after=2, indent_first=False)
para(doc,
     "학습자는 자유 문진을 통해 발병 양상, 통증의 특성, 동반 증상과 위험 신호를 탐색한다. 이어서 선택한 핵심 신체 "
     "진찰을 시행하고, 필요한 검사를 직접 선택한다. 관찰 화면에서는 교육용 영상을 확인한 뒤 발견한 사실과 해석의 "
     "근거를 메모한다. 마지막으로 이 기록을 보고서 형태로 제출하고, 전문가가 설정한 검토 기준과 비교한다.")
figure(doc, "ui-history.png", "그림 2. 자유 입력 기반 가상환자 문진 화면", width_cm=5.9)

para(doc, "3. 교육용 영상의 적용 범위", style_size=10, bold=True, font=SANS,
     space_before=5, space_after=2, indent_first=False)
para(doc,
     "제안 화면의 CT는 사용자 경험과 관찰·기록 흐름 검증을 위한 합성 교육용 이미지이다. 따라서 실제 환자 영상이 "
     "아니며, 임상 판독이나 진단을 위한 자료로 사용하지 않는다. 서비스 적용 단계에서는 비식별화 및 전문가 검수가 "
     "완료된 교육용 영상 자료로 대체하고, 각 영상의 교육 목적과 판독 범위를 명확히 표시한다.")
figure(doc, "ui-imaging.png", "그림 3. 합성 교육용 CT 관찰 및 메모 화면", width_cm=5.9)

para(doc, "4. AI 활용 범위와 안전 설계", style_size=10, bold=True, font=SANS,
     space_before=5, space_after=2, indent_first=False)
para(doc,
     "AI의 역할은 ① 증례 지식베이스에 근거한 가상환자 응답 생성, ② 문진·관찰 기록의 구조화, ③ 학습자가 확인한 "
     "단서와 추가 확인 필요 정보를 비교할 수 있도록 돕는 것으로 한정한다. 반면 최종 진단, 치료 지시, 실제 환자 "
     "상태의 판정은 AI 역할에서 제외한다. 고위험 증례의 내용, 정답 기준 및 교육용 영상은 임상전문가의 검수 절차를 "
     "전제로 한다.")

para(doc, "5. 급성 두통 증례의 교육적 타당성", style_size=10, bold=True, font=SANS,
     space_before=5, space_after=2, indent_first=False)
para(doc,
     "본 초기 증례는 급성 비외상성 두통에서 위험 신호를 놓치지 않고 문진과 진찰의 근거를 설명하는 연습을 목적으로 "
     "한다. 급성 두통 환자에서 연령, 목 통증 또는 강직, 목 굴곡 제한, 의식소실, 발병 양상 등은 추가 평가 필요성을 "
     "검토할 때 활용되는 요소이다[1,2]. 다만 본 연구의 증례는 진료 지침을 대체하거나 실제 진단 결정을 자동화하려는 "
     "목적이 아니며, 교육 목표에 맞춘 제한적 시나리오이다.")

para(doc, "6. 전문가 검토 기반 디브리핑", style_size=10, bold=True, font=SANS,
     space_before=5, space_after=2, indent_first=False)
para(doc,
     "디브리핑은 AI가 학습자의 답을 정오로 채점하는 방식이 아니라, 학습자가 작성한 질문·관찰·검사 선택 이유를 하나의 "
     "학습 보고서로 보여주는 방식이다. 교수자 또는 임상전문가는 보고서에서 확인한 단서, 누락 가능 단서, 검사 선택의 "
     "근거, 불확실성 인식을 검토한다. 이 구조는 결과보다 사고과정을 평가할 수 있게 하며, 반복 학습에서도 학습자별 "
     "피드백 기록을 남길 수 있다.")

# 표 1
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(6)
cap.paragraph_format.space_after = Pt(2)
set_font(cap.add_run("표 1. 단계별 학습자·AI·전문가의 역할 구분"), SANS, 8.5)

rows = [
    ("단계", "학습자 수행", "AI 및 전문가 역할"),
    ("문진", "자유 문장으로 질문", "검수 범위 안의 환자 응답 생성"),
    ("관찰", "진찰·교육용 영상 관찰 기록", "기록 구조화 / 영상 검수"),
    ("보고서", "판단과 근거 제출", "전문가 기준으로 피드백"),
]
table = doc.add_table(rows=0, cols=3)
table.style = "Table Grid"
table.autofit = False
table.allow_autofit = False
# 표가 단 폭을 넘어 세 번째 열이 잘리던 문제 → 레이아웃 고정 + 열 너비 직접 지정
tblPr = table._tbl.tblPr
layout = OxmlElement("w:tblLayout")
layout.set(qn("w:type"), "fixed")
tblPr.append(layout)
tblW = OxmlElement("w:tblW")
tblW.set(qn("w:w"), str(int(7.7 * 567)))
tblW.set(qn("w:type"), "dxa")
tblPr.append(tblW)
widths = [Cm(1.3), Cm(2.8), Cm(3.6)]
for col, w in zip(table.columns, widths):
    col.width = w
for i, row in enumerate(rows):
    cells = table.add_row().cells
    for j, text in enumerate(row):
        cells[j].width = widths[j]
        cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
        set_font(cells[j].paragraphs[0].add_run(text), SANS, 7.5, bold=(i == 0))

heading(doc, "Ⅲ. 기대효과 및 검증 계획")
para(doc,
     "본 서비스는 학습자가 미리 제시된 선택지를 고르는 대신, 필요한 질문을 스스로 구성하고 판단 근거를 설명하도록 "
     "유도한다. 그 결과 위험 단서 탐색, 검사 선택의 정당화, 불확실성 인식 역량을 함께 연습할 수 있다. 향후 의학·"
     "보건계열 학습자를 대상으로 문진 완결성, 근거보고서의 질, 사용성 및 자기효능감을 평가하고, 교수자와 임상전문가의 "
     "내용타당도 검토로 증례 및 피드백 기준을 보완한다.")
para(doc,
     "기존의 가상환자 또는 의료영상 교육 서비스는 대화형 시뮬레이션, 영상 열람, OSCE 평가 중 하나에 초점을 두는 "
     "경우가 많다. 본 시스템은 자유 문진부터 교육용 영상 관찰, 근거보고서 제출, 전문가 검토형 디브리핑을 하나의 "
     "흐름으로 연결하며, 진단 정확도만을 자동 채점하지 않고 어떤 정보를 질문·관찰했고 왜 해당 검사를 선택했는지를 "
     "보고서로 남긴다는 점이 차별점이다. 향후에는 급성 두통 외에도 호흡곤란, 흉통, 감염 의심 등 다양한 고위험·희귀 "
     "증례로 시나리오를 확장할 수 있다.")

heading(doc, "Ⅳ. 결 론")
para(doc,
     "본 연구는 생성형 AI를 진단 도구가 아닌 상호작용 및 학습 근거 구조화 도구로 활용한 임상추론 교육 시뮬레이터를 "
     "제안하였다. 자유 문진, 핵심 진찰, 교육용 영상 관찰, 근거보고서 제출을 하나의 학습 흐름으로 연결함으로써 희귀 "
     "및 응급 증례 학습의 반복성과 설명 가능성을 높일 수 있다. 향후 전문가 검수 증례와 교육용 영상 자료를 연동하고 "
     "교육 효과를 검증하면 실습 수업을 보완하는 교육 매체로 확장될 수 있다.")

heading(doc, "참고문헌")
refs = [
    "[1] National Institute for Health and Care Excellence. Subarachnoid haemorrhage caused by a ruptured "
    "aneurysm: diagnosis and management (NG228). 2022.",
    "[2] Godwin SA, Cherkas DS, Panagos PD, et al. Clinical policy: critical issues in the evaluation and "
    "management of adult patients presenting to the emergency department with acute headache. Ann Emerg Med. "
    "2019;74:e41-e74.",
    "[3] Hoh BL, Ko NU, Amin-Hanjani S, et al. 2023 guideline for the management of patients with aneurysmal "
    "subarachnoid hemorrhage. Stroke. 2023;54:e314-e370.",
    "[4] Generative artificial intelligence for medical image creation in health professions education: a "
    "scoping review. 2026. PubMed PMID: 41663620.",
]
for r in refs:
    p = para(doc, r, style_size=8, space_after=2, indent_first=False)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)

doc.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes")
